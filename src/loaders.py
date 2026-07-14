# ===================================================
# Moss Nexus - Document Loaders
# PDF / Markdown / TXT / DOCX 로더 (LangChain 미사용)
# ===================================================
"""
data/ 디렉토리의 문서를 로드합니다.

설계 원칙:
- 파일 단위로 성공/실패를 추적합니다. 실패는 조용히 넘어가지 않고
  호출자에게 보고되어, 불완전한 코퍼스로 정상 인덱스를 교체하는 사고를 막습니다.
"""

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from loguru import logger

SUPPORTED_EXTENSIONS = (".pdf", ".md", ".txt", ".docx")


@dataclass
class LoadedDocument:
    """로드된 문서 (PDF는 페이지 단위로 분리됨)"""
    text: str
    source: str                 # data/ 기준 상대 경로
    filename: str
    file_hash: str              # 원본 파일 전체의 sha256
    page: int | None = None     # PDF 페이지 번호 (1부터)
    metadata: dict = field(default_factory=dict)


@dataclass
class LoadError:
    """로드 실패 기록"""
    source: str
    error: str


def _read_text_file(path: Path) -> str:
    """UTF-8 우선, 한국어 레거시 인코딩(cp949) 폴백으로 텍스트를 읽습니다."""
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp949")


def _load_pdf(path: Path) -> list[tuple[str, int | None]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append((text, i))
    return pages


def _load_docx(path: Path) -> list[tuple[str, int | None]]:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [("\n".join(parts), None)] if parts else []


def _load_plain(path: Path) -> list[tuple[str, int | None]]:
    text = _read_text_file(path)
    return [(text, None)] if text.strip() else []


_LOADERS = {
    ".pdf": _load_pdf,
    ".md": _load_plain,
    ".txt": _load_plain,
    ".docx": _load_docx,
}


def load_documents(data_path: Path) -> tuple[list[LoadedDocument], list[LoadError]]:
    """
    data_path 아래의 지원 문서를 모두 로드합니다.

    Returns:
        (documents, errors): 로드된 문서 리스트와 파일 단위 실패 기록.
        errors가 비어 있지 않으면 호출자는 색인 교체 여부를 신중히 결정해야 합니다.
    """
    documents: list[LoadedDocument] = []
    errors: list[LoadError] = []

    if not data_path.exists():
        logger.warning(f"데이터 디렉토리가 존재하지 않습니다: {data_path}")
        return documents, errors

    files = sorted(
        p for p in data_path.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )
    logger.info(f"문서 로드 시작: {data_path} ({len(files)}개 파일)")

    for path in files:
        rel = str(path.relative_to(data_path))
        try:
            file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            loader = _LOADERS[path.suffix.lower()]
            parts = loader(path)
            if not parts:
                # 스캔 이미지 PDF, 빈 파일 등 — 조용히 빠뜨리지 않고 실패로 보고
                raise ValueError("텍스트를 추출하지 못했습니다 (스캔 이미지 PDF 또는 빈 파일?)")
            for text, page in parts:
                documents.append(LoadedDocument(
                    text=text,
                    source=rel,
                    filename=path.name,
                    file_hash=file_hash,
                    page=page,
                ))
        except Exception as e:
            logger.error(f"문서 로드 실패: {rel} ({type(e).__name__})")
            errors.append(LoadError(source=rel, error=f"{type(e).__name__}: {e}"))

    logger.info(f"문서 로드 완료: 성공 {len(documents)}건, 실패 {len(errors)}건")
    return documents, errors
